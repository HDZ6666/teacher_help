import io
import math
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from exceptions.exception import ServiceException


class ImportRowError(Exception):
    """
    单行校验失败（只在导入服务内部使用，汇总为错误行回报）
    """


class TeachImportUtil:
    """
    教学域 Excel 导入公共方法：模板生成、读取、单元格解析、错误行汇总
    """

    MAX_ROWS = 2000
    MAX_FILE_SIZE = 5 * 1024 * 1024
    PHONE_PATTERN = re.compile(r'^1\d{10}$')

    @classmethod
    def build_template(
        cls,
        sheet_name: str,
        headers: list[str],
        notes: list[str],
        options: dict[str, list[str]] | None = None,
        example: list | None = None,
    ) -> bytes:
        """
        生成导入模板：第一行表头（* 为必填，红色），可选第二行示例；“填写说明”页写规则；下拉选项放在隐藏页 options
        """
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name
        header_fill = PatternFill(start_color='DDEBF7', end_color='DDEBF7', fill_type='solid')
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')
            cell.font = Font(bold=True, color='C00000' if header.startswith('*') else '000000')
            ws.column_dimensions[get_column_letter(col)].width = max(14, len(header) * 2 + 4)
        if example:
            for col, value in enumerate(example, 1):
                ws.cell(row=2, column=col, value=value)
        note_ws = wb.create_sheet('填写说明')
        note_ws.column_dimensions['A'].width = 110
        for index, note in enumerate(notes, 1):
            note_ws.cell(row=index, column=1, value=note)
        if options:
            opt_ws = wb.create_sheet('options')
            for opt_col, (header, values) in enumerate(options.items(), 1):
                if header not in headers or not values:
                    continue
                letter = get_column_letter(opt_col)
                for opt_row, value in enumerate(values, 1):
                    opt_ws.cell(row=opt_row, column=opt_col, value=value)
                dv = DataValidation(
                    type='list', formula1=f'options!${letter}$1:${letter}${len(values)}', allow_blank=True
                )
                target = get_column_letter(headers.index(header) + 1)
                dv.add(f'{target}2:{target}{cls.MAX_ROWS + 1}')
                ws.add_data_validation(dv)
            opt_ws.sheet_state = 'hidden'
        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()

    @classmethod
    async def read_rows(cls, file, required_headers: list[str]):
        """
        读取第一个工作表，返回 [(Excel行号, {列名: 文本})]；列名去掉 * 与空白
        """
        file_bytes = await file.read()
        if not file_bytes:
            raise ServiceException(message='导入文件为空')
        if len(file_bytes) > cls.MAX_FILE_SIZE:
            raise ServiceException(message='导入文件不能超过5MB')
        try:
            data_frame = pd.read_excel(io.BytesIO(file_bytes), sheet_name=0, dtype=object)
        except Exception as e:
            raise ServiceException(message=f'读取Excel失败，请使用下载的模板：{e}')
        data_frame.columns = [str(col).replace('*', '').strip() for col in data_frame.columns]
        missing = [header for header in required_headers if header not in data_frame.columns]
        if missing:
            raise ServiceException(message=f'导入失败，模板缺少列：{"、".join(missing)}，请重新下载模板')
        rows = []
        for index, raw in data_frame.iterrows():
            values = {col: cls.to_text(raw.get(col)) for col in data_frame.columns}
            if not any(values.values()):
                continue
            rows.append((int(index) + 2, values))
        if not rows:
            raise ServiceException(message='导入文件没有数据行')
        if len(rows) > cls.MAX_ROWS:
            raise ServiceException(message=f'单次最多导入{cls.MAX_ROWS}行，请拆分文件')
        return rows

    @classmethod
    def to_text(cls, value):
        if value is None:
            return None
        if isinstance(value, float):
            if math.isnan(value):
                return None
            if value.is_integer():
                return str(int(value))
            return str(value)
        if isinstance(value, (datetime, pd.Timestamp)):
            return value.strftime('%Y-%m-%d')
        if isinstance(value, date):
            return value.isoformat()
        text = str(value).strip()
        return text or None

    @classmethod
    def required(cls, values: dict, name: str):
        value = values.get(name)
        if not value:
            raise ImportRowError(f'{name}不能为空')
        return value

    @classmethod
    def phone(cls, values: dict, name: str = '手机号'):
        value = cls.required(values, name)
        if not cls.PHONE_PATTERN.match(value):
            raise ImportRowError(f'{name}必须是11位手机号')
        return value

    @classmethod
    def parse_date(cls, values: dict, name: str, required: bool = False):
        value = values.get(name)
        if not value:
            if required:
                raise ImportRowError(f'{name}不能为空')
            return None
        text = value.split(' ')[0].replace('/', '-').replace('.', '-')
        try:
            return datetime.strptime(text, '%Y-%m-%d').date()
        except ValueError:
            raise ImportRowError(f'{name}格式不正确，应为 YYYY-MM-DD')

    @classmethod
    def parse_int(cls, values: dict, name: str, required: bool = False, default: int = 0):
        value = values.get(name)
        if not value:
            if required:
                raise ImportRowError(f'{name}不能为空（没有请填0）')
            return default
        try:
            number = Decimal(value)
        except InvalidOperation:
            raise ImportRowError(f'{name}必须是数字')
        if number < 0 or number != number.to_integral_value():
            raise ImportRowError(f'{name}必须是大于等于0的整数（当前课程账户按整数计数）')
        return int(number)

    @classmethod
    def parse_money(cls, values: dict, name: str, required: bool = False):
        value = values.get(name)
        if not value:
            if required:
                raise ImportRowError(f'{name}不能为空')
            return Decimal('0')
        try:
            number = Decimal(value.replace(',', ''))
        except InvalidOperation:
            raise ImportRowError(f'{name}必须是数字')
        if number < 0:
            raise ImportRowError(f'{name}不能小于0')
        return number.quantize(Decimal('0.01'))

    @classmethod
    def parse_decimal(cls, values: dict, name: str, default: Decimal):
        value = values.get(name)
        if not value:
            return default
        try:
            number = Decimal(value)
        except InvalidOperation:
            raise ImportRowError(f'{name}必须是数字')
        if number <= 0:
            raise ImportRowError(f'{name}必须大于0')
        return number.quantize(Decimal('0.01'))

    @classmethod
    def result(cls, total: int, errors: list[dict], success_count: int = 0, extra: dict | None = None):
        data = {
            'success': not errors,
            'total': total,
            'successCount': success_count,
            'errorCount': len(errors),
            'errors': errors[:500],
        }
        if extra:
            data.update(extra)
        return data
