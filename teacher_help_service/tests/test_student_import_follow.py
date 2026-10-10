"""
学员导入（3 个模板）、跟进记录、收据、班级导入/升班/结业（P3）
"""

import io
from datetime import date
from types import SimpleNamespace

from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from tests.helpers import raises_service
from module_teach.entity.do.teach_class_do import TeachClass, TeachClassStudent
from module_teach.entity.do.teach_course_do import TeachCourse
from module_teach.entity.do.teach_course_price_do import TeachCoursePrice
from module_teach.entity.do.teach_enrollment_order_do import TeachEnrollmentOrder, TeachStudentCourseAccount
from module_teach.entity.do.teach_student_do import TeachStudent
from module_teach.entity.vo.teach_student_account_vo import (
    ClassBatchIdsModel,
    ClassPromoteModel,
    TeachStudentFollowModel,
    TeachStudentFollowPageQueryModel,
)
from module_teach.service.teach_class_batch_service import TeachClassBatchService
from module_teach.service.teach_student_account_service import TeachStudentAccountService
from module_teach.service.teach_student_follow_service import TeachStudentFollowService
from module_teach.service.teach_student_import_service import TeachStudentImportService


class FakeUpload:
    def __init__(self, data: bytes):
        self.data = data

    async def read(self):
        return self.data


def make_xlsx(headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return FakeUpload(buffer.getvalue())


async def make_course(factory, name, charge_type='class'):
    course = await factory.add(TeachCourse(course_name=name, status=1, del_flag=0))
    await factory.add(TeachCoursePrice(course_id=course.id, charge_type=charge_type, price_name='单价'))
    return course


async def test_templates_have_headers(db, factory):
    await make_course(factory, '小学数学')
    await db.commit()
    for kind, headers in (
        ('student', TeachStudentImportService.STUDENT_HEADERS),
        ('lesson', TeachStudentImportService.LESSON_HEADERS),
        ('month', TeachStudentImportService.MONTH_HEADERS),
    ):
        data = await TeachStudentImportService.get_template_services(db, kind)
        ws = load_workbook(io.BytesIO(data)).active
        assert [cell.value for cell in ws[1]] == headers
    with raises_service('导入类型'):
        await TeachStudentImportService.get_template_services(db, 'bad')
    data = await TeachClassBatchService.get_import_template_services(db)
    assert load_workbook(io.BytesIO(data)).active['A1'].value == '*班级名称'


async def test_student_import_all_or_nothing(db, factory):
    headers = TeachStudentImportService.STUDENT_HEADERS
    bad = make_xlsx(
        headers,
        [
            ['13800009999', '新学员', '男', '2016-01-02', '实验小学', '二年级', '新家长', '', '', ''],
            ['123', '错手机', '', '', '', '', '', '', '', ''],
            ['13800009999', '新学员', '', '', '', '', '', '', '', ''],
        ],
    )
    result = await TeachStudentImportService.import_services(db, 'student', bad, 'tester')
    assert not result['success'] and {e['row'] for e in result['errors']} == {3, 4}
    assert (await db.execute(select(TeachStudent).where(TeachStudent.student_name == '新学员'))).first() is None

    good = make_xlsx(
        headers,
        [
            ['13800009999', '新学员', '男', '2016-01-02', '实验小学', '二年级', '新家长', '', '', ''],
            ['13800009999', '新学员妹妹', '女', '', '', '', '', '', '', ''],
        ],
    )
    result = await TeachStudentImportService.import_services(db, 'student', good, 'tester')
    assert result['success'] and result['createdStudents'] == 2
    students = (await db.execute(select(TeachStudent).where(TeachStudent.student_name.like('新学员%')))).scalars().all()
    assert len({s.parent_id for s in students}) == 1 and {s.gender for s in students} == {'1', '2'}

    again = await TeachStudentImportService.import_services(db, 'student', good, 'tester')
    assert not again['success'] and '已存在' in again['errors'][0]['message']


async def test_enrollment_imports_create_accounts(db, factory):
    lesson_course = await make_course(factory, '小学数学')
    month_course = await make_course(factory, '托管课', charge_type='month')
    class_obj = await factory.add(TeachClass(class_name='数学1班', course_id=lesson_course.id, course_name='小学数学'))
    await db.commit()

    lesson = make_xlsx(
        TeachStudentImportService.LESSON_HEADERS,
        [
            [
                '13800008888',
                '麦小小',
                '男',
                '小学数学',
                '数学1班',
                30,
                5,
                2,
                '',
                3000,
                2500,
                '微信',
                '2026-10-01',
                '2027-10-01',
                '',
                '',
            ],
        ],
    )
    result = await TeachStudentImportService.import_services(db, 'lesson', lesson, 'tester')
    assert result['success'], result
    account = (
        (
            await db.execute(
                select(TeachStudentCourseAccount).where(TeachStudentCourseAccount.course_id == lesson_course.id)
            )
        )
        .scalars()
        .first()
    )
    assert (
        account.purchased_quantity,
        account.gift_quantity,
        account.consumed_quantity,
        account.remaining_quantity,
    ) == (30, 5, 2, 33)
    assert account.class_id == class_obj.id
    order = await db.get(TeachEnrollmentOrder, account.order_id)
    assert order.order_type == 'import' and order.paid_amount == 2500
    member = (
        (await db.execute(select(TeachClassStudent).where(TeachClassStudent.class_id == class_obj.id)))
        .scalars()
        .first()
    )
    assert member.course_account_id == account.id

    receipt = await TeachStudentAccountService.get_receipt_services(db, order.id)
    assert receipt['debtAmount'] == 500 and receipt['items'][0]['itemName'] == '小学数学'

    month = make_xlsx(
        TeachStudentImportService.MONTH_HEADERS,
        [
            [
                '13800008888',
                '麦小小',
                '',
                '托管课',
                '',
                '2026-10-01',
                '2027-01-31',
                '2027-02-28',
                6000,
                6000,
                '',
                '',
                '',
                '',
            ],
        ],
    )
    result = await TeachStudentImportService.import_services(db, 'month', month, 'tester')
    assert result['success'] and result['createdStudents'] == 0, result
    account = (
        (
            await db.execute(
                select(TeachStudentCourseAccount).where(TeachStudentCourseAccount.course_id == month_course.id)
            )
        )
        .scalars()
        .first()
    )
    assert (account.purchased_quantity, account.gift_quantity, account.valid_end_date) == (4, 1, date(2027, 2, 28))

    wrong = make_xlsx(
        TeachStudentImportService.LESSON_HEADERS,
        [
            ['13800008888', '麦小小', '', '托管课', '', 1, 0, 2, '', 1, 1, '', '', '', '', ''],
            ['13800008888', '麦小小', '', '小学数学', '', 1.5, 0, 0, '', 1, 1, '', '', '', '', ''],
        ],
    )
    result = await TeachStudentImportService.import_services(db, 'lesson', wrong, 'tester')
    messages = [e['message'] for e in result['errors']]
    assert not result['success'] and '按课时' in messages[0] and '整数' in messages[1]


async def test_follow_crud(db, factory):
    student = await factory.student('跟进学员')
    await db.commit()
    user = SimpleNamespace(user_id=1, user_name='admin', nick_name='管理员')
    form = TeachStudentFollowModel(
        studentId=student.id, followType='phone', content='电话沟通续费', nextFollowDate=date(2099, 1, 1)
    )
    result = await TeachStudentFollowService.save_follow_services(db, form, user)
    follow_id = result.result['id']
    form = TeachStudentFollowModel(id=follow_id, studentId=student.id, followType='wechat', content='改为微信')
    await TeachStudentFollowService.save_follow_services(db, form, user)
    page = await TeachStudentFollowService.get_follow_list_services(
        db, TeachStudentFollowPageQueryModel(studentId=student.id)
    )
    assert page.total == 1 and page.rows[0]['followTypeName'] == '微信' and page.rows[0]['followUserName'] == '管理员'
    with raises_service('不能早于'):
        await TeachStudentFollowService.save_follow_services(
            db,
            TeachStudentFollowModel(
                studentId=student.id, followType='visit', content='x', nextFollowDate=date(2000, 1, 1)
            ),
            user,
        )
    await TeachStudentFollowService.delete_follow_services(db, str(follow_id), 'admin')
    page = await TeachStudentFollowService.get_follow_list_services(
        db, TeachStudentFollowPageQueryModel(studentId=student.id)
    )
    assert page.total == 0


async def test_class_import_promote_graduate(db, factory):
    course = await make_course(factory, '数学启蒙')
    await db.commit()
    upload = make_xlsx(
        TeachClassBatchService.IMPORT_HEADERS,
        [
            ['启蒙一班', '数学启蒙', 20, '不可超额', 5, '1号教室', '', 1.5, ''],
            ['启蒙二班', '不存在课程', '', '可超额', '', '', '', '', ''],
        ],
    )
    result = await TeachClassBatchService.import_services(db, upload, 'tester')
    assert not result['success'] and result['errors'][0]['row'] == 3
    upload = make_xlsx(
        TeachClassBatchService.IMPORT_HEADERS,
        [
            ['启蒙一班', '数学启蒙', 20, '不可超额', 5, '1号教室', '', 1.5, ''],
        ],
    )
    result = await TeachClassBatchService.import_services(db, upload, 'tester')
    assert result['success']
    imported = (await db.execute(select(TeachClass).where(TeachClass.class_name == '启蒙一班'))).scalars().first()
    assert imported.allow_over_capacity == 0 and imported.max_students == 20

    s1 = await factory.student('升班甲')
    s2 = await factory.student('升班乙')
    a1 = await factory.course_account(s1, 10, course_id=course.id)
    a2 = await factory.course_account(s2, 10, course_id=course.id)
    old = await factory.class_with_students([(s1, a1), (s2, a2)])
    old.course_id = course.id
    await db.commit()
    result = await TeachClassBatchService.promote_services(
        db,
        ClassPromoteModel(items=[{'sourceClassId': old.id, 'className': '启蒙进阶班'}], graduateSource=True),
        'tester',
    )
    new_id = result.result[0]['newClassId']
    assert result.result[0]['promotedStudents'] == 2 and result.result[0]['unboundStudents'] == []
    members = (await db.execute(select(TeachClassStudent).where(TeachClassStudent.class_id == new_id))).scalars().all()
    assert {m.course_account_id for m in members} == {a1.id, a2.id}
    old = await db.get(TeachClass, old.id, populate_existing=True)
    assert old.enroll_status == 3 and old.current_students == 0

    result = await TeachClassBatchService.graduate_services(db, ClassBatchIdsModel(classIds=[new_id]), 'tester')
    graduated = await db.get(TeachClass, new_id, populate_existing=True)
    assert graduated.enroll_status == 3 and graduated.current_students == 0 and result.result[0]['removedStudents'] == 2
    assert (await db.get(TeachStudentCourseAccount, a1.id, populate_existing=True)).remaining_quantity == 10
    with raises_service('已结业'):
        await TeachClassBatchService.graduate_services(db, ClassBatchIdsModel(classIds=[new_id]), 'tester')
